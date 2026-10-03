import numpy as np
import pandas as pd


def weighted_percentile(x, w) -> pd.Series:
    """가중 백분위(0~1). 동점은 중간순위. x가 NaN이거나 w<=0인 행은 분포 계산에서 제외하고,
    NaN이 아닌 x에는 보간으로 값을 준다."""
    x = pd.Series(x, dtype=float)
    w = pd.Series(w, index=x.index, dtype=float).clip(lower=0)
    out = pd.Series(np.nan, index=x.index)
    ok = x.notna() & (w > 0)
    if ok.sum() == 0:
        return out
    g = w[ok].groupby(x[ok]).sum().sort_index()
    pct = (g.cumsum() - g + 0.5 * g) / g.sum()
    has = x.notna()
    out[has] = np.interp(x[has], g.index.to_numpy(), pct.to_numpy())
    return out


def weighted_median(x, w) -> float:
    d = pd.DataFrame({"x": x, "w": w}).dropna()
    d = d[d.w > 0].sort_values("x")
    if d.empty:
        return np.nan
    return d.x[d.w.cumsum() >= d.w.sum() / 2].iloc[0]


def grade_balanced(score, w, labels=("낮음", "중간", "높음"), share_range=None) -> pd.Series:
    """동점을 쪼개지 않는 등급. 경계는 서로 다른 점수값 사이에만 두고, 누적 가중 비율이
    1/k, 2/k에 가장 가까운 위치를 앞에서부터 고른다. 같은 응답을 한 사람은 항상 같은 등급이다.
    share_range=(lo, hi)를 주면, 등급 비율이 범위를 벗어날 때(3등급 한정) 가장 큰 동점 덩어리를
    가운데 등급에 고정하고 그보다 낮은 값/높은 값을 양쪽 등급으로 둔다."""
    s = pd.Series(score, dtype=float)
    w = pd.Series(w, index=s.index, dtype=float).clip(lower=0)
    out = pd.Series(np.nan, index=s.index, dtype=object)
    ok = s.notna() & (w > 0)
    if not ok.any():
        return out
    g = w[ok].groupby(s[ok]).sum().sort_index()
    vals, cum = g.index.to_numpy(), (g.cumsum() / g.sum()).to_numpy()
    k, cuts, lo = len(labels), [], 0
    for j in range(1, k):
        hi = len(vals) - (k - j)                 # 남은 등급마다 값이 하나 이상 남도록
        if hi <= lo:
            break                                # 서로 다른 값이 등급 수보다 적다
        i = lo + int(np.argmin(np.abs(cum[lo:hi] - j / k)))
        cuts.append(vals[i])
        lo = i + 1
    if share_range and k == 3 and len(vals) >= 3:
        lab = np.searchsorted(cuts, vals, side="left")
        sh = np.bincount(lab, weights=g.to_numpy(), minlength=k) / g.sum()
        if ((sh < share_range[0]) | (sh > share_range[1])).any():
            p = int(np.argmax(g.to_numpy()))     # 가장 큰 동점 덩어리 -> 가운데 등급
            cuts = [vals[p - 1] if p > 0 else -np.inf, vals[p]]
    has = s.notna()
    out[has] = [labels[b] for b in np.searchsorted(cuts, s[has].to_numpy(), side="left")]
    return out


def grade_shares(grade, w) -> dict:
    """등급별 가중 비율."""
    t = pd.Series(w, index=grade.index).groupby(grade).sum()
    return {str(k): round(float(v), 3) for k, v in (t / t.sum()).items()}


def allocate(weights: pd.Series, n: int) -> pd.Series:
    """최대 나머지 방식으로 n명을 칸별로 배분."""
    raw = weights / weights.sum() * n
    q = np.floor(raw).astype(int)
    rest = int(n - q.sum())
    if rest > 0:
        order = (raw - q).sort_values(ascending=False, kind="stable").index[:rest]
        q.loc[order] += 1
    return q


def codebook_lookup(var, year=2024, path=None):
    """예: codebook_lookup('p__income') -> (설명, [(코드, 라벨, 해당 연도 가중 %), ...])"""
    from openpyxl import load_workbook

    from . import config as C
    wb = load_workbook(path or C.CODEBOOK, read_only=True)
    for sn in wb.sheetnames[2:]:
        rows = list(wb[sn].iter_rows(values_only=True))
        if len(rows) < 4 or year not in rows[2]:
            continue
        c = list(rows[2]).index(year)
        cur, desc, out = None, None, []
        for r in rows[3:]:
            if len(r) > 1 and r[1]:
                cur = r[1]
                if cur == var:
                    desc = str(r[2]).replace("\n", " ")
            if cur == var and len(r) > c:
                out.append((r[3], r[4], r[c]))
        if out:
            return desc, out
    return None, []
